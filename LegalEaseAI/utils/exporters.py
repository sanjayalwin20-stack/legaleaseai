from io import BytesIO
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF

from PIL import Image

from utils.text_utils import (
    parse_terms,
    sanitize_text
)


def format_txt(
    text: str
) -> bytes:

    return sanitize_text(
        text
    ).encode(
        "utf-8"
    )


def add_logo_to_docx(
    doc: Document,
    logo_bytes: Optional[bytes]
):

    if not logo_bytes:
        return

    try:

        stream = BytesIO(
            logo_bytes
        )

        Image.open(
            stream
        ).verify()

        stream.seek(0)

        paragraph = doc.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run()

        run.add_picture(
            stream,
            width=Inches(1.4)
        )

    except Exception:

        # Logo is optional.
        # Invalid logo should not stop
        # document generation.

        return


def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    logo_bytes: Optional[bytes] = None
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(
        0.7
    )

    section.bottom_margin = Inches(
        0.7
    )

    section.left_margin = Inches(
        0.85
    )

    section.right_margin = Inches(
        0.85
    )

    styles = document.styles

    styles["Normal"].font.name = (
        "Times New Roman"
    )

    styles["Normal"].font.size = Pt(
        11
    )

    add_logo_to_docx(
        document,
        logo_bytes
    )

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        sanitize_text(
            doc_type
        ).upper()
    )

    title_run.bold = True

    title_run.font.name = (
        "Times New Roman"
    )

    title_run.font.size = Pt(
        16
    )

    clean_text = sanitize_text(
        text
    )

    for raw_line in clean_text.split(
        "\n"
    ):

        line = raw_line.strip()

        if not line:
            continue

        paragraph = (
            document.add_paragraph()
        )

        is_heading = (
            len(line) < 100
            and (
                line.isupper()
                or line.lower().startswith(
                    (
                        "section ",
                        "article ",
                        "clause "
                    )
                )
            )
        )

        run = paragraph.add_run(
            line
        )

        run.font.name = (
            "Times New Roman"
        )

        if is_heading:

            run.bold = True
            run.font.size = Pt(
                12
            )

        else:

            run.font.size = Pt(
                11
            )

            paragraph.paragraph_format.space_after = Pt(
                6
            )

    parsed_terms = parse_terms(
        terms
    )

    if parsed_terms:

        document.add_page_break()

        heading = (
            document.add_paragraph()
        )

        heading_run = heading.add_run(
            "Key Terms"
        )

        heading_run.bold = True

        heading_run.font.name = (
            "Times New Roman"
        )

        heading_run.font.size = Pt(
            13
        )

        table = document.add_table(
            rows=1,
            cols=2
        )

        table.style = "Table Grid"

        table.rows[0].cells[0].text = (
            "No."
        )

        table.rows[0].cells[1].text = (
            "Term"
        )

        for index, term in enumerate(
            parsed_terms,
            start=1
        ):

            cells = (
                table.add_row().cells
            )

            cells[0].text = str(
                index
            )

            cells[1].text = term

    footer = (
        section.footer.paragraphs[0]
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - AI-generated draft. "
        "Review with a qualified legal professional."
    )

    footer_run.font.size = Pt(
        8
    )

    output = BytesIO()

    document.save(
        output
    )

    return output.getvalue()


class LegalEasePDF(FPDF):

    def __init__(
        self,
        doc_type: str,
        logo_path: Optional[str] = None
    ):

        super().__init__()

        self.doc_type = doc_type

        self.logo_path = logo_path

        self.set_auto_page_break(
            auto=True,
            margin=18
        )

    def header(self):

        if (
            self.logo_path
            and Path(
                self.logo_path
            ).exists()
        ):

            try:

                self.image(
                    self.logo_path,
                    x=95,
                    y=8,
                    w=20
                )

                self.ln(18)

            except Exception:

                pass

        self.set_font(
            "Helvetica",
            "B",
            12
        )

        self.cell(
            0,
            8,
            sanitize_text(
                self.doc_type
            ),
            align="C"
        )

        self.ln(8)

    def footer(self):

        self.set_y(-14)

        self.set_font(
            "Helvetica",
            "",
            7
        )

        self.cell(
            0,
            8,
            (
                "LegalEase - AI-generated draft - "
                "Professional legal review recommended"
            ),
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str,
    logo_bytes: Optional[bytes] = None
) -> bytes:

    temporary_logo = None

    try:

        if logo_bytes:

            temporary_logo = (
                "/tmp/legalease_logo.png"
            )

            image = Image.open(
                BytesIO(logo_bytes)
            )

            image.convert(
                "RGB"
            ).save(
                temporary_logo,
                "PNG"
            )

        pdf = LegalEasePDF(
            doc_type,
            temporary_logo
        )

        pdf.add_page()

        clean_text = sanitize_text(
            text
        )

        for raw_line in clean_text.split(
            "\n"
        ):

            line = raw_line.strip()

            if not line:

                pdf.ln(3)

                continue

            is_heading = (
                len(line) < 100
                and (
                    line.isupper()
                    or line.lower().startswith(
                        (
                            "section ",
                            "article ",
                            "clause "
                        )
                    )
                )
            )

            pdf.set_font(
                "Helvetica",
                "B" if is_heading else "",
                11 if is_heading else 10
            )

            pdf.multi_cell(
                0,
                6,
                line
            )

            pdf.ln(1)

        return bytes(
            pdf.output()
        )

    finally:

        if temporary_logo:

            try:

                Path(
                    temporary_logo
                ).unlink(
                    missing_ok=True
                )

            except Exception:

                pass